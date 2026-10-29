"""Read an answer document back with the common Python readers and compare what it shows with its request body.

    python read_back.py <request.json> <document.docx|xlsx|pptx>

Needs python-docx, openpyxl and python-pptx (for example `pixi exec --spec python-docx --spec openpyxl --spec
python-pptx python read_back.py ...`). Word: the paragraphs in order and every table cell by row and column. Excel:
each section's sheet, a table's header and rows from A1 with numbers and truth values typed, then the Provenance sheet.
PowerPoint: each slide's title and every table cell. Exits 1 with the first difference; prints only counts otherwise.
"""
import json
import re
import sys


def cell_text(value):
    if value is None:
        return ""
    if value is True or value is False:
        return "true" if value else "false"
    return value if isinstance(value, str) else json.dumps(value)


def fail(message):
    print("DIFFERENT:", message)
    sys.exit(1)


def word(request, path):
    from docx import Document

    doc = Document(path)
    expected = [request["title"]]
    tables = []
    for section in request["sections"]:
        expected.append(section["heading"])
        expected += section.get("paragraphs", [])
        for block in ("table", "figure"):
            if section.get(block):
                if block == "table":
                    table = section["table"]
                    tables.append([table["columns"]] + [[cell_text(v) for v in row] for row in table["rows"]])
                expected += [t for t in (section[block].get("caption", ""), section[block]["ref"]) if t]
    expected.append("Provenance")
    expected += [f"{e['label']}: {e['value']}" for e in request["provenance"]]
    shown = [p.text for p in doc.paragraphs if p.text or p.style.name != "Figure"]
    shown = [t for t in shown if t != ""]
    if shown != [t for t in expected if t != ""]:
        fail(f"paragraphs: {len(shown)} shown, {len(expected)} sent")
    grids = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    if grids != tables:
        fail("a table's cells")
    return f"{len(shown)} paragraphs, {len(grids)} tables"


def excel_text(value):
    """A cell's text as Excel shows it: openpyxl (3.1) leaves SpreadsheetML's `_xHHHH_` escapes in inline strings,
    where Excel decodes them (`_x005F_` is the underscore the writer escapes a literal `_xHHHH_` with)."""
    if not isinstance(value, str):
        return value
    return re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m.group(1), 16)), value)


def workbook(request, path):
    from openpyxl import load_workbook

    book = load_workbook(path)
    sheets = book.worksheets
    if len(sheets) != len(request["sections"]) + 1:
        fail("sheet count")
    for sheet, section in zip(sheets, request["sections"]):
        table = section.get("table")
        if table:
            header = [c.value for c in sheet[1][: len(table["columns"])]]
            if header != table["columns"]:
                fail(f"{sheet.title}: header")
            for r, row in enumerate(table["rows"], 2):
                shown = [excel_text(sheet.cell(r, c + 1).value) for c in range(len(row))]
                sent = [None if v == "" else v for v in row]
                if shown != sent:
                    fail(f"{sheet.title}: row {r}")
        else:
            texts = [request["title"], section["heading"], *section.get("paragraphs", [])]
            if section.get("figure"):
                texts += [section["figure"].get("caption", ""), section["figure"]["ref"]]
            shown = [excel_text(sheet.cell(r + 1, 1).value) or "" for r in range(len(texts))]
            if shown != texts:
                fail(f"{sheet.title}: text rows")
    provenance = sheets[-1]
    rows = [[provenance.cell(r, c).value for c in (1, 2)] for r in range(2, len(request["provenance"]) + 2)]
    if rows != [[e["label"], e["value"]] for e in request["provenance"]]:
        fail("provenance")
    return f"{len(sheets)} sheets"


def slides(request, path):
    from pptx import Presentation

    deck = Presentation(path)
    if len(deck.slides) != len(request["sections"]) + 2:
        fail("slide count")
    titles = [s.shapes[0].text_frame.text for s in deck.slides]
    if titles != [request["title"], *(s["heading"] for s in request["sections"]), "Provenance"]:
        fail("slide titles")
    for slide, section in zip(list(deck.slides)[1:], request["sections"]):
        table = section.get("table")
        grids = [[[c.text for c in row.cells] for row in s.table.rows] for s in slide.shapes if s.has_table]
        if table and grids != [[table["columns"]] + [[cell_text(v) for v in row] for row in table["rows"]]]:
            fail(f"{section['heading']}: table cells")
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        for paragraph in section.get("paragraphs", []):
            if not any(paragraph.replace("\n", "\x0b") in t or paragraph in t for t in texts):
                fail(f"{section['heading']}: a paragraph")
    return f"{len(deck.slides)} slides"


def main():
    request = json.load(open(sys.argv[1]))
    path = sys.argv[2]
    reader = {"docx": word, "xlsx": workbook, "pptx": slides}[path.rsplit(".", 1)[1]]
    print(request["kind"], "matches:", reader(request, path))


if __name__ == "__main__":
    main()
